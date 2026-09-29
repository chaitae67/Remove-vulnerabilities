#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""클라우드 결과보고서 양식(빈 xlsx) 생성기 — CSP별로 한 번 만들어 커밋한다.

서버(DBMS) 결과보고서와 동일한 5시트 구성. 항목코드·세부항목·진단기준·영역을
cloud_check/*_items 에서 박아 넣고 판정/근거(F/G)만 비운다.
 - 2-2 요약 판정은 3-1 참조 수식, 영역별 양호율은 COUNTIF 수식(자동 계산)
 - 2-1 그래프: 3D막대·3D원형·레이더 3차트
 - 판정 색은 조건부서식(취약=빨강, N/A=회색, 인터뷰=파랑)
이후 server_report.py(cloud 스펙)가 표지/진단대상/3-1 F·G 만 채운다.

  python build_cloud_templates.py [aws|azure|gcp|naver ...]
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart3D, PieChart3D, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, FormulaRule

PROVIDERS = {
    "aws":   ("aws_items",   "AWS",   "SK Shieldus 2024 클라우드 보안가이드 (AWS)",   "보고서_양식_AWS.xlsx"),
    "azure": ("azure_items", "Azure", "SK Shieldus 2024 클라우드 보안가이드 (Azure)", "보고서_양식_Azure.xlsx"),
    "gcp":   ("gcp_items",   "GCP",   "SK Shieldus 2024 클라우드 보안가이드 (GCP)",   "보고서_양식_GCP.xlsx"),
    "naver": ("ncp_items",   "Naver", "네이버 클라우드 플랫폼 보안 가이드",           "보고서_양식_Naver.xlsx"),
}

GRAY_HDR = PatternFill("solid", fgColor="FFD9D9D9")
GRAY_LBL = PatternFill("solid", fgColor="FFBFBFBF")
ORANGE = PatternFill("solid", fgColor="FFFCD5B5")
THIN = Side(style="thin", color="FFB2B2B2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
DET = "3-1. 진단 결과"
SUM = "2-2. 요약 진단결과"
GRAPH = "2-1. 요약결과(그래프)"
FIRST = 6


def _ck(c):
    import re
    m = re.match(r"(\d+)\.(\d+)$", str(c))
    if m:
        return (0, int(m.group(1)), int(m.group(2)))
    m = re.match(r"([A-Za-z]+)-?(\d+)$", str(c))
    return (1, m.group(1), int(m.group(2))) if m else (2, str(c), 0)


def _h(ws, r, c, t):
    x = ws.cell(r, c, t); x.fill = GRAY_HDR; x.font = BOLD; x.alignment = CEN; x.border = BORDER
    return x


def build(provider):
    mod, label, guide, outname = PROVIDERS[provider]
    items = __import__("cloud_check." + mod, fromlist=["ITEMS"]).ITEMS
    codes = sorted(items.keys(), key=_ck)
    n = len(codes)
    last = FIRST + n - 1
    # 영역별 연속 행 범위
    areas, area_rows = [], {}
    r = FIRST
    for c in codes:
        a = items[c].get("area", "기타")
        if a not in area_rows:
            areas.append(a); area_rows[a] = [r, r]
        area_rows[a][1] = r
        r += 1
    row_of = {c: FIRST + i for i, c in enumerate(codes)}

    wb = openpyxl.Workbook()
    cov = wb.active; cov.title = "0. 표지"
    tgt = wb.create_sheet("1. 진단 대상")
    wg = wb.create_sheet(GRAPH)
    sm = wb.create_sheet(SUM)
    dt = wb.create_sheet(DET)

    # ---- 0. 표지 ----
    cov.sheet_view.showGridLines = False
    for col, w in zip("BCDEFGHIJKL", (13, 13, 13, 13, 13, 13, 13, 13, 4, 11, 22)):
        cov.column_dimensions[col].width = w
    for i, (k, v) in enumerate([("문서번호", "XXXXX-VA-2026XXX"), ("작성자", "취약점진단팀"),
                                ("보안등급", "Confidential"), ("Ver", "ver 1.0")]):
        a = cov.cell(3 + i, 11, k); a.fill = GRAY_LBL; a.font = BOLD; a.alignment = CEN; a.border = BORDER
        b = cov.cell(3 + i, 12, v); b.alignment = CEN; b.border = BORDER
    for rng, text, font in (("B11:I11", '"진단대상" 취약점 진단', Font(bold=True, size=16)),
                            ("B13:I13", f"클라우드({label}) 진단 상세결과", Font(bold=True, size=24)),
                            ("B18:I18", "", Font(bold=True, size=14)),
                            ("B21:I21", guide, Font(size=10, color="FF808080"))):
        cov.merge_cells(rng)
        cc = cov[rng.split(":")[0]]; cc.value = text; cc.font = font
        cc.alignment = Alignment(horizontal="center", vertical="center")

    # ---- 1. 진단 대상 ----
    tgt.cell(1, 2, f"  ※ 진단 대상 리스트 - 클라우드({label}) 1대").font = Font(bold=True, size=13)
    _h(tgt, 2, 2, "순번"); _h(tgt, 2, 3, "진단 대상"); tgt.merge_cells("C2:F2"); _h(tgt, 2, 7, "비고")
    for c, t in ((3, "계정/구독"), (4, "ID"), (5, "리전/버전"), (6, "용도")):
        _h(tgt, 3, c, t)
    b4 = tgt.cell(4, 2, label); b4.fill = ORANGE; b4.font = BOLD; b4.alignment = CEN; b4.border = BORDER
    for c in range(2, 7):
        cc = tgt.cell(5, c); cc.border = BORDER; cc.alignment = CEN
    for col, w in zip("BCDEFG", (8, 22, 22, 20, 24, 6)):
        tgt.column_dimensions[col].width = w

    # ---- 3-1. 진단 결과(상세) ----
    dt.cell(2, 2, f"클라우드({label}) 취약점 진단 상세결과({n}항목)").font = Font(bold=True, size=14)
    for c, t in ((2, "진단항목"), (3, "항목코드"), (4, "세부 진단항목"), (5, "진단기준"), (6, "판정"), (7, "상세 내용 / 근거")):
        _h(dt, 3, c, t)
    for c in codes:
        m = items[c]; rr = row_of[c]
        dt.cell(rr, 3, c); dt.cell(rr, 4, m.get("title", "")); dt.cell(rr, 5, m.get("crit", ""))
        for col in range(2, 8):
            cc = dt.cell(rr, col); cc.border = BORDER; cc.alignment = CEN if col in (3, 6) else LEFT
    for a in areas:
        s, e = area_rows[a]
        dt.cell(s, 2, a)
        if e > s:
            dt.merge_cells(start_row=s, start_column=2, end_row=e, end_column=2)
    for col, w in zip("BCDEFG", (16, 11, 26, 46, 8, 60)):
        dt.column_dimensions[col].width = w
    fc = f"F{FIRST}:F{last}"
    dt.conditional_formatting.add(fc, CellIsRule(operator="equal", formula=['"취약"'], font=Font(color="FFFF0000", bold=True)))
    dt.conditional_formatting.add(fc, CellIsRule(operator="equal", formula=['"N/A"'], font=Font(color="FF808080")))
    dt.conditional_formatting.add(fc, FormulaRule(formula=[f'ISNUMBER(SEARCH("인터뷰",F{FIRST}))'], font=Font(color="FF0070C0", bold=True)))
    dt.freeze_panes = "A4"

    # ---- 2-2. 요약 진단결과 ----
    sm.cell(2, 2, f"클라우드({label}) 취약점 진단 요약결과({n}항목)").font = Font(bold=True, size=14)
    for c, t in ((2, "진단항목"), (3, "항목코드"), (4, "세부 진단항목"), (5, "중요도"), (6, "진단결과")):
        _h(sm, 3, c, t)
    _h(sm, 3, 8, "진단 영역"); _h(sm, 3, 9, "양호율")
    for c in codes:
        m = items[c]; rr = row_of[c]
        sm.cell(rr, 3, c); sm.cell(rr, 4, m.get("title", "")); sm.cell(rr, 5, m.get("imp", ""))
        sm.cell(rr, 6, f"='{DET}'!F{rr}")
        for col in range(2, 7):
            cc = sm.cell(rr, col); cc.border = BORDER; cc.alignment = CEN if col in (3, 5, 6) else LEFT
    for a in areas:
        s, e = area_rows[a]
        sm.cell(s, 2, a)
        if e > s:
            sm.merge_cells(start_row=s, start_column=2, end_row=e, end_column=2)
    for i, a in enumerate(areas):
        s, e = area_rows[a]; rr = FIRST + i
        sm.cell(rr, 8, a).border = BORDER; sm.cell(rr, 8).alignment = LEFT
        rng = f"F{s}:F{e}"
        cell = sm.cell(rr, 9, f'=IFERROR(COUNTIF({rng},"양호")/(COUNTA({rng})-COUNTIF({rng},"N/A")-COUNTIF({rng},"인터뷰 필요")),1)')
        cell.number_format = "0.0%"; cell.border = BORDER; cell.alignment = CEN
    for col, w in zip("BCDEFHI", (16, 11, 30, 8, 12, 20, 10)):
        sm.column_dimensions[col].width = w
    fc = f"F{FIRST}:F{last}"
    sm.conditional_formatting.add(fc, CellIsRule(operator="equal", formula=['"취약"'], font=Font(color="FFFF0000", bold=True)))
    sm.conditional_formatting.add(fc, CellIsRule(operator="equal", formula=['"N/A"'], font=Font(color="FF808080")))
    sm.conditional_formatting.add(fc, FormulaRule(formula=[f'ISNUMBER(SEARCH("인터뷰",F{FIRST}))'], font=Font(color="FF0070C0", bold=True)))

    # ---- 2-1. 요약결과(그래프) ----
    wg.sheet_view.showGridLines = False
    wg.cell(1, 1, f"클라우드({label}) 진단 요약").font = Font(bold=True, size=13)
    _h(wg, 3, 1, "진단 영역"); _h(wg, 3, 2, "양호율")
    for i, a in enumerate(areas):
        wg.cell(4 + i, 1, a).border = BORDER
        cc = wg.cell(4 + i, 2, f"='{SUM}'!I{FIRST + i}"); cc.number_format = "0.0%"; cc.border = BORDER; cc.alignment = CEN
    a_last = 4 + len(areas) - 1
    _h(wg, 3, 4, "구분"); _h(wg, 3, 5, "건수")
    for i, lab in enumerate(["양호", "취약", "인터뷰 필요", "N/A"]):
        wg.cell(4 + i, 4, lab).border = BORDER
        wg.cell(4 + i, 5, f'=COUNTIF(\'{SUM}\'!$F${FIRST}:$F${last},"{lab}")').border = BORDER
    wg.column_dimensions["A"].width = 20; wg.column_dimensions["D"].width = 12

    bar = BarChart3D(); bar.title = "영역별 양호율"; bar.type = "col"; bar.legend = None; bar.height = 8; bar.width = 13
    bar.add_data(Reference(wg, min_col=2, min_row=3, max_row=a_last), titles_from_data=True)
    bar.set_categories(Reference(wg, min_col=1, min_row=4, max_row=a_last))
    bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True
    wg.add_chart(bar, "A10")
    pie = PieChart3D(); pie.title = "진단 결과 분포"; pie.height = 8; pie.width = 13
    pie.add_data(Reference(wg, min_col=5, min_row=3, max_row=7), titles_from_data=True)
    pie.set_categories(Reference(wg, min_col=4, min_row=4, max_row=7))
    pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
    wg.add_chart(pie, "J10")
    radar = RadarChart(); radar.type = "filled"; radar.title = "영역별 양호율(레이더)"; radar.style = 26
    radar.height = 8; radar.width = 13; radar.legend = None
    radar.add_data(Reference(wg, min_col=2, min_row=3, max_row=a_last), titles_from_data=True)
    radar.set_categories(Reference(wg, min_col=1, min_row=4, max_row=a_last))
    wg.add_chart(radar, "A28")

    try:
        wb.calc_properties.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(outname)
    print(f"양식 저장: {outname}  ({n}항목, 영역 {len(areas)})")


def main():
    for p in (sys.argv[1:] or list(PROVIDERS)):
        build(p)


if __name__ == "__main__":
    main()
