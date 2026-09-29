#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""인프라(웹서버/DBMS/클라우드) 진단 결과 → 서버 보고서와 동일한 5시트 xlsx.

구성(서버 리눅스/윈도우 결과보고서 양식과 통일):
  0. 표지 / 1. 진단 대상 / 2-1. 요약결과(그래프, 3차트) / 2-2. 요약 진단결과 / 3-1. 진단 결과(상세)
  - 2-1: 영역별 점수(막대 3D) · 등급 분포(원형 3D) · 영역별 양호율(레이더)
  - 판정 색: 취약=빨강, 인터뷰 필요=파랑
  - 진단기준(3-1 E열): 클라우드는 *_items 에서 자동, 웹/DB 는 아래 CRIT 내장.
"""
import datetime

MANUAL_LABEL = "인터뷰 필요"
REPORT_STATUS = {
    "양호": "양호", "취약": "취약", "N/A": "양호",
    "수동확인": MANUAL_LABEL, MANUAL_LABEL: MANUAL_LABEL,
}
TARGET_LABEL = {
    "web": "웹서버", "nginx": "웹서버(Nginx)", "iis": "웹서버(IIS)", "tomcat": "웹서버(Tomcat)",
    "dbms": "DBMS(Oracle)", "oracle": "DBMS(Oracle)",
    "aws": "클라우드(AWS)", "azure": "클라우드(Azure)", "gcp": "클라우드(GCP)", "naver": "클라우드(Naver)",
}
GUIDE = {
    "web": "KISA 주통기 상세가이드(WEB) / SK Shieldus 웹 보안가이드",
    "dbms": "KISA 주통기 상세가이드(DBMS) / SK Shieldus DB 보안가이드",
    "aws": "SK Shieldus 2024 클라우드 보안가이드 (AWS)",
    "azure": "SK Shieldus 2024 클라우드 보안가이드 (Azure)",
    "gcp": "SK Shieldus 2024 클라우드 보안가이드 (GCP)",
    "naver": "네이버 클라우드 플랫폼 보안 가이드",
}
CLOUD = {"aws", "azure", "gcp", "naver"}

# 진단 영역(도메인) 매핑
_WEB_AREA = {}
for _c in range(1, 27):
    _k = "WEB-%02d" % _c
    _WEB_AREA[_k] = ("1. 계정 관리" if _c <= 3 else "2. 서비스 관리" if _c <= 18
                     else "3. 보안 설정" if _c <= 23 else "4. 패치 및 로그 관리")
_DB_AREA = {}
for _c in range(1, 27):
    _k = "D-%02d" % _c
    _DB_AREA[_k] = ("1. 계정 관리" if _c <= 9 else "2. 접근 관리" if _c <= 16
                    else "3. 옵션 관리" if _c <= 24 else "4. 패치 관리")

# 진단기준(양호/취약) — 웹/DB (공식 결과보고서 문구 축약)
CRIT_WEB = {
    "WEB-01": "양호 - 관리자 페이지 미사용 또는 계정명이 기본값이 아님 / 취약 - 기본 계정명 또는 추측 쉬운 계정명 사용",
    "WEB-02": "양호 - 관리자 비밀번호가 암호화되어 있거나 유추 어려움 / 취약 - 미암호화 또는 유추 쉬운 비밀번호",
    "WEB-03": "양호 - 비밀번호 파일 권한 600 이하 / 취약 - 600 초과",
    "WEB-04": "양호 - 디렉터리 리스팅 미설정 / 취약 - 디렉터리 리스팅 설정",
    "WEB-05": "양호 - CGI 미사용 또는 실행 디렉터리 제한 / 취약 - CGI 사용 + 미제한",
    "WEB-06": "양호 - 상위 디렉터리 접근 기능 제거 / 취약 - 미제거",
    "WEB-07": "양호 - 불필요한 기본 파일·디렉터리 없음 / 취약 - 존재",
    "WEB-08": "양호 - 업로드/다운로드 용량 제한 / 취약 - 미제한",
    "WEB-09": "양호 - 최소권한 별도 계정으로 구동 / 취약 - 관리자 권한 계정으로 구동",
    "WEB-10": "양호 - 불필요한 Proxy 설정 제한 / 취약 - 미제한",
    "WEB-11": "양호 - 업무영역 분리 경로 + 불필요 경로 없음 / 취약 - 미분리 또는 불필요 경로 존재",
    "WEB-12": "양호 - 심볼릭 링크·alias 등 링크 미허용 / 취약 - 허용",
    "WEB-13": "양호 - DB 연결 파일 접근 제한 + 불필요 매핑 제거 / 취약 - 미제한 또는 미제거",
    "WEB-14": "양호 - 주요 파일·디렉터리 불필요 권한 없음 / 취약 - 불필요 접근 권한 부여",
    "WEB-15": "양호 - 불필요한 스크립트 매핑 없음 / 취약 - 존재",
    "WEB-16": "양호 - 응답 헤더에 서버 정보 미노출 / 취약 - 노출",
    "WEB-17": "양호 - 불필요한 가상 디렉터리 없음 / 취약 - 존재",
    "WEB-18": "양호 - WebDAV 비활성화 / 취약 - 활성화",
    "WEB-19": "양호 - SSI 비활성화 / 취약 - 활성화",
    "WEB-20": "양호 - SSL/TLS 활성화 / 취약 - 비활성화",
    "WEB-21": "양호 - HTTP→HTTPS 리디렉션 활성 / 취약 - 비활성",
    "WEB-22": "양호 - 에러 페이지 별도 지정 / 취약 - 미지정 또는 중요정보 노출",
    "WEB-23": "양호 - LDAP 안전한 다이제스트 알고리즘 사용 / 취약 - 미사용",
    "WEB-24": "양호 - 별도 업로드 경로 + 일반 사용자 접근권한 없음 / 취약 - 미분리 또는 권한 부여",
    "WEB-25": "양호 - 최신 보안 패치 + 패치 정책 수립 / 취약 - 미적용",
    "WEB-26": "양호 - 로그 디렉터리·파일에 일반 사용자 접근 없음 / 취약 - 접근 권한 있음",
}
CRIT_D = {
    "D-01": "양호 - 기본 계정 초기 비밀번호 변경 또는 잠금 / 취약 - 미변경·미잠금",
    "D-02": "양호 - 불필요한 계정 없음 / 취약 - 불필요(인가되지 않은/테스트) 계정 존재",
    "D-03": "양호 - 비밀번호 사용기간·복잡도 정책 적용 / 취약 - 미적용",
    "D-04": "양호 - 관리자 권한을 필요한 계정/그룹에만 부여 / 취약 - 불필요 계정에 부여",
    "D-05": "양호 - 비밀번호 재사용 제한 설정 / 취약 - 미설정",
    "D-06": "양호 - 사용자/응용별 개별 계정 부여(불필요 계정 없음) / 취약 - 불필요 계정 존재",
    "D-07": "양호 - root 아닌 계정/권한으로 구동 / 취약 - root 계정/권한으로 구동",
    "D-08": "양호 - SHA-256 이상 해시 알고리즘 사용 / 취약 - 미만",
    "D-09": "양호 - 로그인 실패 잠금값 설정 / 취약 - 미설정",
    "D-10": "양호 - 지정 IP에서만 접근 가능하도록 제한 / 취약 - 미제한",
    "D-11": "양호 - 시스템 테이블 DBA만 접근 / 취약 - 일반 계정 접근 가능",
    "D-12": "양호 - 리스너 비밀번호 설정 / 취약 - 미설정",
    "D-13": "양호 - 불필요한 ODBC/OLE-DB 미설치 / 취약 - 설치",
    "D-14": "양호 - 주요 파일 일반 사용자 수정권한 제거 / 취약 - 미제거",
    "D-15": "양호 - 리스너 설정파일 관리자 권한 + 파라미터 변경 제한 / 취약 - 일반 사용자 권한/변경 가능",
    "D-16": "양호 - Windows 인증 모드 + sa 비활성 / 취약 - 혼합 인증 + sa 약한 암호",
    "D-17": "양호 - Audit Table 관리자 계정만 접근 / 취약 - 일반 계정 접근",
    "D-18": "양호 - DBA 계정 Role 이 Public 아님 / 취약 - Public 설정",
    "D-19": "양호 - OS_ROLES/REMOTE_OS_AUTHENT/REMOTE_OS_ROLES 가 FALSE / 취약 - TRUE",
    "D-20": "양호 - Object Owner 가 관리자 계정으로 제한 / 취약 - 일반 사용자 소유 존재",
    "D-21": "양호 - WITH GRANT OPTION 이 ROLE 로 설정 / 취약 - 미설정",
    "D-22": "양호 - RESOURCE_LIMIT 가 TRUE / 취약 - FALSE",
    "D-23": "양호 - xp_cmdshell 비활성 또는 조건 충족 / 취약 - 활성 + 조건 미충족",
    "D-24": "양호 - Registry Procedure 가 DBA 외 guest/public 미부여 / 취약 - 부여",
    "D-25": "양호 - 보안 패치 적용 버전 사용 / 취약 - 미적용",
    "D-26": "양호 - 감사 로그 저장 정책 수립 + 적용 / 취약 - 미수립·미적용",
}


def _meta_for(target):
    """(area_map, crit_map, items_dict) 반환. 클라우드는 *_items 모듈에서."""
    t = (target or "").lower()
    if t in CLOUD:
        mod = {"aws": "aws_items", "azure": "azure_items", "gcp": "gcp_items", "naver": "ncp_items"}[t]
        import importlib
        try:
            items = importlib.import_module("cloud_check." + mod).ITEMS
        except Exception:
            try:
                items = importlib.import_module(mod).ITEMS
            except Exception:
                items = {}
        area = {c: v.get("area", "기타") for c, v in items.items()}
        crit = {c: v.get("crit", "") for c, v in items.items()}
        return area, crit, items
    if t in ("dbms", "oracle"):
        return _DB_AREA, CRIT_D, {}
    return _WEB_AREA, CRIT_WEB, {}


def _code_key(c):
    import re
    m = re.match(r"(\d+)\.(\d+)$", str(c))
    if m:
        return (0, int(m.group(1)), int(m.group(2)), "")
    m = re.match(r"([A-Za-z]+)-?(\d+)$", str(c))
    if m:
        return (1, m.group(1), int(m.group(2)), "")
    return (2, str(c), 0, "")


def build_report(target, host, osver, results, out_path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import RadarChart, BarChart3D, PieChart3D, Reference

    target = (target or "").lower()
    label = TARGET_LABEL.get(target, target.upper())
    area_map, crit_map, items = _meta_for(target)
    today = datetime.date.today().strftime("%Y. %m. %d.")

    thin = Side(style="thin", color="FFBFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="FF44546A")
    area_fill = PatternFill("solid", fgColor="FFF2F2F2")
    hdr_font = Font(bold=True, color="FFFFFFFF")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    red = Font(color="FFFF0000", bold=True)
    blue = Font(color="FF0070C0", bold=True)

    def verdict_of(r):
        return REPORT_STATUS.get(r.get("status", ""), r.get("status", ""))

    def area_of(code):
        return area_map.get(code, items.get(code, {}).get("area", "기타"))

    def title_of(r):
        c = r.get("code", "")
        return r.get("title") or items.get(c, {}).get("title", "")

    def imp_of(r):
        c = r.get("code", "")
        return r.get("importance") or items.get(c, {}).get("imp", "")

    def crit_of(code):
        return crit_map.get(code, items.get(code, {}).get("crit", ""))

    results = sorted(results, key=lambda r: _code_key(r.get("code", "")))
    areas = []
    for r in results:
        a = area_of(r.get("code", ""))
        if a not in areas:
            areas.append(a)

    def hdr(ws, row, col, text):
        c = ws.cell(row, col, text)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        c.border = border
        return c

    wb = openpyxl.Workbook()

    # ================= 0. 표지 =================
    ws = wb.active
    ws.title = "0. 표지"
    ws.sheet_view.showGridLines = False
    for i, (k, v) in enumerate([("문서번호", "XXXXX-VA-2026XXX"), ("작성자", "취약점진단팀"),
                                ("보안등급", "Confidential"), ("Ver", "ver 1.0")]):
        r = 3 + i
        a = ws.cell(r, 11, k); a.fill = area_fill; a.font = bold; a.alignment = center; a.border = border
        b = ws.cell(r, 12, v); b.alignment = center; b.border = border
    ws.cell(11, 2, f'"{host}" 취약점 진단').font = Font(bold=True, size=18)
    ws.cell(13, 2, f"{label} 진단 상세결과").font = Font(bold=True, size=22, color="FF1F4E79")
    ws.cell(18, 2, today).font = Font(bold=True, size=12)
    ws.cell(21, 2, GUIDE.get(target, "")).font = Font(size=10, color="FF808080")
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 22

    # ================= 1. 진단 대상 =================
    ws = wb.create_sheet("1. 진단 대상")
    ws.cell(1, 2, f"  ※ 진단 대상 - {label}").font = Font(bold=True, size=13)
    for j, h in enumerate(["구분", "대상 / 계정", "버전 / 리전", "비고"]):
        hdr(ws, 3, 2 + j, h)
    row = [label, host, osver or "-", GUIDE.get(target, "")]
    for j, v in enumerate(row):
        c = ws.cell(4, 2 + j, v); c.alignment = left; c.border = border
    for col, w in (("B", 16), ("C", 32), ("D", 26), ("E", 44)):
        ws.column_dimensions[col].width = w

    # ---- 집계 ----
    tot_good = sum(1 for r in results if verdict_of(r) == "양호")
    tot_vuln = sum(1 for r in results if verdict_of(r) == "취약")
    tot_man = sum(1 for r in results if verdict_of(r) == MANUAL_LABEL)
    tot_na = len(results) - tot_good - tot_vuln - tot_man
    denom = tot_good + tot_vuln
    total_rate = (tot_good / denom) if denom else 1.0
    area_rate = {}
    for a in areas:
        g = sum(1 for r in results if area_of(r.get("code", "")) == a and verdict_of(r) == "양호")
        v = sum(1 for r in results if area_of(r.get("code", "")) == a and verdict_of(r) == "취약")
        area_rate[a] = (g / (g + v)) if (g + v) else 1.0

    # ================= 2-1. 요약결과(그래프) =================
    wg = wb.create_sheet("2-1. 요약결과(그래프)")
    wg.sheet_view.showGridLines = False
    wg.cell(1, 1, f"{label} 진단 요약 (전체 보안 적용율 {total_rate*100:.1f}%)").font = Font(bold=True, size=13)
    # 영역별 점수표 (막대/레이더 소스) : A3 헤더, A4..
    hdr(wg, 3, 1, "진단 영역"); hdr(wg, 3, 2, "양호율")
    ar0 = 4
    for i, a in enumerate(areas):
        wg.cell(ar0 + i, 1, a).border = border
        c = wg.cell(ar0 + i, 2, round(area_rate[a], 3)); c.number_format = "0.0%"; c.border = border; c.alignment = center
    ar_last = ar0 + len(areas) - 1
    # 등급 분포표 (원형 소스) : D3 헤더
    hdr(wg, 3, 4, "구분"); hdr(wg, 3, 5, "건수")
    dist = [("양호", tot_good), ("취약", tot_vuln), (MANUAL_LABEL, tot_man), ("N/A", tot_na)]
    for i, (k, v) in enumerate(dist):
        wg.cell(4 + i, 4, k).border = border
        wg.cell(4 + i, 5, v).border = border
    wg.column_dimensions["A"].width = 20
    wg.column_dimensions["D"].width = 12

    # 막대 3D — 영역별 양호율
    bar = BarChart3D(); bar.title = "영역별 양호율"; bar.style = 10; bar.height = 8; bar.width = 13
    bar.add_data(Reference(wg, min_col=2, min_row=3, max_row=ar_last), titles_from_data=True)
    bar.set_categories(Reference(wg, min_col=1, min_row=ar0, max_row=ar_last))
    bar.legend = None
    wg.add_chart(bar, "A10")
    # 원형 3D — 등급 분포
    pie = PieChart3D(); pie.title = "진단 결과 분포"; pie.height = 8; pie.width = 13
    pie.add_data(Reference(wg, min_col=5, min_row=3, max_row=7), titles_from_data=True)
    pie.set_categories(Reference(wg, min_col=4, min_row=4, max_row=7))
    wg.add_chart(pie, "J10")
    # 레이더 — 영역별 양호율
    radar = RadarChart(); radar.type = "filled"; radar.title = "영역별 양호율(레이더)"; radar.style = 26
    radar.height = 8; radar.width = 13
    radar.add_data(Reference(wg, min_col=2, min_row=3, max_row=ar_last), titles_from_data=True)
    radar.set_categories(Reference(wg, min_col=1, min_row=ar0, max_row=ar_last))
    radar.legend = None
    wg.add_chart(radar, "A28")

    # ================= 2-2. 요약 진단결과 =================
    ws = wb.create_sheet("2-2. 요약 진단결과")
    ws.cell(2, 2, f"{label} 취약점 진단 요약결과({len(results)}항목)").font = Font(bold=True, size=12)
    for j, h in enumerate(["진단항목", "항목코드", "세부 진단항목", "중요도", "진단결과"]):
        hdr(ws, 3, 2 + j, h)
    r = 4
    for a in areas:
        first = r
        for res in [x for x in results if area_of(x.get("code", "")) == a]:
            v = verdict_of(res)
            ws.cell(r, 2, a if r == first else None)
            ws.cell(r, 3, res.get("code", ""))
            ws.cell(r, 4, title_of(res))
            ws.cell(r, 5, imp_of(res))
            vc = ws.cell(r, 6, v)
            if v == "취약":
                vc.font = red
            elif v == MANUAL_LABEL:
                vc.font = blue
            for col in range(2, 7):
                cc = ws.cell(r, col); cc.border = border
                cc.alignment = center if col in (3, 5, 6) else left
            r += 1
    for col, w in zip("BCDEF", (16, 10, 34, 8, 12)):
        ws.column_dimensions[col].width = w

    # ================= 3-1. 진단 결과(상세) =================
    ws = wb.create_sheet("3-1. 진단 결과")
    ws.cell(2, 2, f"{label} 취약점 진단 상세결과({len(results)}항목)").font = Font(bold=True, size=12)
    for j, h in enumerate(["진단항목", "항목코드", "세부 진단항목", "진단기준", "진단결과", "상세 내용 / 근거"]):
        hdr(ws, 3, 2 + j, h)
    r = 4
    for a in areas:
        first = r
        for res in [x for x in results if area_of(x.get("code", "")) == a]:
            code = res.get("code", "")
            v = verdict_of(res)
            ws.cell(r, 2, a if r == first else None)
            ws.cell(r, 3, code)
            ws.cell(r, 4, title_of(res))
            ws.cell(r, 5, crit_of(code))
            vc = ws.cell(r, 6, v)
            if v == "취약":
                vc.font = red
            elif v == MANUAL_LABEL:
                vc.font = blue
            ev = " / ".join(res.get("evidence", []))
            if res.get("resources"):
                ev = (ev + "  [리소스: " + " / ".join(res.get("resources", [])) + "]") if ev else ev
            ws.cell(r, 7, ev[:32000])
            for col in range(2, 8):
                cc = ws.cell(r, col); cc.border = border
                cc.alignment = center if col in (3, 6) else left
            r += 1
    for col, w in zip("BCDEFG", (16, 10, 28, 50, 12, 60)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    wb.save(out_path)
    return out_path
