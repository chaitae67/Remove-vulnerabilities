#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""클라우드 취약점 진단 — GUI 없이 실행하는 단독 스크립트(CLI).

★ 클라우드 쉘에서 키 없이 그대로:
    (AWS CloudShell)      python cloud_scan.py aws
    (GCP Cloud Shell)     python cloud_scan.py gcp
    (Azure Cloud Shell)   python cloud_scan.py azure --subscription-id <ID>
  → 쉘에 이미 로그인된 자격(ambient: CloudShell/역할/az login/ADC)을 자동 사용한다.
    GCP 프로젝트는 Cloud Shell 환경변수에서 자동 인식(없으면 --project).

키를 직접 줄 수도 있다(쉘 밖·CI 등):
    python cloud_scan.py aws   --access-key AKIA... --secret-key ... --region ap-northeast-2
    python cloud_scan.py azure --tenant-id .. --client-id .. --client-secret .. --subscription-id ..
    python cloud_scan.py gcp   --sa-key sa.json --project my-proj
    python cloud_scan.py naver --access-key .. --secret-key .. --region KR   (또는 환경변수 NCP_ACCESS_KEY/NCP_SECRET_KEY)

네이버(NCP)는 쉘 기본자격이 없어 키가 필요하며, 안 주면 실행 중 물어본다.
진단이 끝나면 콘솔에 요약을 찍고, 해당 CSP 보고서 양식(.xlsx)에 채워 저장한다.

의존성: 해당 CSP SDK (aws=boto3, azure=azure-identity/mgmt, gcp=google-api-python-client,
naver=표준 라이브러리) + openpyxl. tkinter/GUI 불필요.
클라우드 쉘에는 boto3/gcloud SDK 가 대개 미리 깔려 있어 바로 실행된다.
"""
import argparse
import datetime
import getpass
import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)

# 콘솔 인코딩이 UTF-8 이 아니어도(윈도우 cp949 등) 한글/기호 출력이 깨지거나 죽지 않게
for _s in (getattr(sys, "stdout", None), getattr(sys, "stderr", None)):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import cloud_check  # noqa: E402  (같은 폴더의 패키지)

MANUAL_LABEL = "인터뷰 필요"
REPORT_STATUS = {
    "양호": "양호", "취약": "취약", "N/A": "N/A",
    "수동확인": MANUAL_LABEL, MANUAL_LABEL: MANUAL_LABEL,
}

# CSP(run 이 돌려주는 os 라벨) → 양식파일 ASCII 접미사(한글 파일명 인코딩에 의존하지 않음)
TEMPLATES = {"AWS": "_AWS.xlsx", "AZURE": "_Azure.xlsx",
             "GCP": "_GCP.xlsx", "NAVER": "_Naver.xlsx"}
FIRST_ROW, CODE_COL, RESULT_COL, DETAIL_COL, RES_COL = 4, 2, 7, 8, 9


# ------------------------------------------------------------------ 자격증명
def _ask(label, secret=False, default=""):
    # 비대화형(파이프/CI)에서는 입력을 기다리지 않고 기본값 사용 → 행 방지
    if not sys.stdin or not sys.stdin.isatty():
        return default
    prompt = f"{label}{f' [{default}]' if default else ''}: "
    try:
        val = getpass.getpass(prompt) if secret else input(prompt)
    except (EOFError, KeyboardInterrupt):
        val = ""
    return (val or default).strip()


def _env(*names):
    for n in names:
        v = os.environ.get(n)
        if v:
            return v
    return ""


def collect_creds(provider, args):
    """자격증명을 모은다. AWS/Azure/GCP 는 클라우드 쉘/로그인 자격(ambient)을 그대로 쓰도록
    인자·환경변수만 참고하고 프롬프트하지 않는다. NCP 만 없으면 물어본다."""
    p = provider.lower()
    if p == "aws":
        # 키를 안 주면 boto3 기본 체인(CloudShell/EC2 역할/env/~/.aws)을 사용
        ak = args.access_key or _env("AWS_ACCESS_KEY_ID")
        sk = args.secret_key or _env("AWS_SECRET_ACCESS_KEY")
        creds = {
            "access_key": ak, "secret_key": sk,
            "session_token": args.session_token or _env("AWS_SESSION_TOKEN"),
            "profile": args.profile or "",
            "region": args.region or _env("AWS_REGION", "AWS_DEFAULT_REGION") or "ap-northeast-2",
        }
        creds["mode"] = "key" if ak else ("profile" if creds["profile"] else "env")
        return creds
    if p == "azure":
        # 비우면 DefaultAzureCredential(az login / Cloud Shell / 관리 ID) 사용
        return {
            "tenant_id": args.tenant_id or _env("AZURE_TENANT_ID"),
            "client_id": args.client_id or _env("AZURE_CLIENT_ID"),
            "client_secret": args.client_secret or _env("AZURE_CLIENT_SECRET"),
            "subscription_id": args.subscription_id or _env("AZURE_SUBSCRIPTION_ID"),
        }
    if p == "gcp":
        # 비우면 ADC. Cloud Shell 은 프로젝트가 환경변수로 들어있음
        return {
            "sa_key_path": args.sa_key or _env("GOOGLE_APPLICATION_CREDENTIALS"),
            "project_id": args.project or _env("GOOGLE_CLOUD_PROJECT", "DEVSHELL_PROJECT_ID",
                                               "GCLOUD_PROJECT", "GCP_PROJECT"),
        }
    if p == "naver":
        # NCP 는 ambient 자격이 없어 키가 필요 → 인자/환경변수, 없으면 물어본다
        ak = args.access_key or _env("NCP_ACCESS_KEY", "NCP_ACCESS_KEY_ID")
        sk = args.secret_key or _env("NCP_SECRET_KEY", "NCP_SECRET_ACCESS_KEY")
        if not ak:
            ak = _ask("NCP Access Key ID")
        if not sk:
            sk = _ask("NCP Secret Key", secret=True)
        return {"access_key": ak, "secret_key": sk,
                "region": args.region or _env("NCP_REGION") or "KR"}
    raise SystemExit(f"알 수 없는 provider: {provider}")


# ------------------------------------------------------------------ 결과 저장
def _base_name(os_label, out_path):
    """저장 파일 기본 경로(확장자 제외). 한글/특수문자 없는 ASCII 이름 → 쉘 다운로드 호환."""
    if out_path:
        root, ext = os.path.splitext(out_path)
        return root if ext.lower() in (".xlsx", ".csv") else out_path
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    return os.path.join(os.getcwd(), f"cloud_{(os_label or 'result').upper()}_{stamp}")


def save_results(provider, os_label, host, results, out_path=None, want_xlsx=True):
    """CSV(항상, 설치 불필요) + 다중시트 보고서 XLSX(openpyxl 있으면) 저장, 경로 목록 반환."""
    base = _base_name(os_label, out_path)
    saved = []
    # 1) CSV — 표준 라이브러리, 한글 깨짐 방지(utf-8-sig). 설치 없이도 항상 저장.
    import csv
    csv_path = base + ".csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["항목코드", "진단항목", "중요도", "진단결과", "상세", "리소스"])
        for r in results:
            verdict = REPORT_STATUS.get(r.get("status", ""), r.get("status", ""))
            w.writerow([r.get("code", ""), r.get("title", ""), r.get("importance", ""),
                        verdict, " / ".join(r.get("evidence", [])),
                        " / ".join(r.get("resources", []))])
    saved.append(csv_path)
    # 2) 보고서 XLSX — 공식 CSP 양식(5시트+3차트) 채우기 우선, 안 되면 코드 생성 폴백
    if want_xlsx:
        xlsx = base + ".xlsx"
        if not _fill_cloud_template(provider, host, os_label, results, xlsx):
            try:
                import cloud_check.report as _rep
                saved.append(_rep.build_report(provider, host, results, xlsx))
            except ImportError:
                pass  # openpyxl 미설치 → CSV 로 충분
            except Exception as e:  # noqa: BLE001
                print(f"    (엑셀 보고서 생성 실패: {type(e).__name__}: {e} → CSV 로 저장됨)")
        else:
            saved.append(xlsx)
    return saved


def _fill_cloud_template(provider, host, os_label, results, out_path):
    """공식 CSP 양식(보고서_양식_<CSP>.xlsx)을 찾아 server_report 로 채운다. 성공 시 True."""
    suffix = {"aws": "_AWS.xlsx", "azure": "_Azure.xlsx",
              "gcp": "_GCP.xlsx", "naver": "_Naver.xlsx"}.get(provider)
    if not suffix:
        return False
    try:
        import glob
        import lxml.etree  # noqa: F401  (server_report 의존성)
        import server_report
    except Exception:
        return False
    def _has_cover(path):
        # 새 5시트 양식만 사용(옛 1시트 양식 배제): '0. 표지' 시트가 있어야 함
        try:
            import zipfile
            import re as _re
            with zipfile.ZipFile(path) as z:
                wb = z.read("xl/workbook.xml").decode("utf-8", "ignore")
            return "0. 표지" in "".join(_re.findall(r'name="([^"]+)"', wb))
        except Exception:
            return False

    tpl = None
    seen = set()
    dirs = [os.path.dirname(os.path.abspath(__file__))] + [p for p in sys.path if p]
    for d in dirs:
        try:
            hits = [f for f in glob.glob(os.path.join(d, "*.xlsx")) if f.endswith(suffix)]
        except Exception:
            hits = []
        for f in hits:
            rp = os.path.realpath(f)
            if rp in seen:
                continue
            seen.add(rp)
            if _has_cover(f):
                tpl = f
                break
        if tpl:
            break
    if not tpl:
        return False
    sv = {"host": host or provider.upper(), "ip": "-", "osver": os_label or provider.upper(),
          "role": "클라우드 계정", "results": results}
    try:
        server_report.fill_report("cloud", [sv], tpl, out_path)
        return True
    except Exception as e:  # noqa: BLE001
        print(f"    (양식 채우기 실패: {type(e).__name__}: {e} → 코드 생성으로 폴백)")
        return False


def _find_template(key):
    """SCRIPT_DIR 에서 ASCII 접미사(_AWS.xlsx 등)로 양식 파일을 찾는다(한글 프리픽스 무관)."""
    suffix = TEMPLATES.get(key)
    if not suffix:
        return None
    try:
        for name in os.listdir(SCRIPT_DIR):
            if name.endswith(suffix):
                return os.path.join(SCRIPT_DIR, name)
    except OSError:
        pass
    return None


def _pick_sheet(wb):
    """항목코드가 B열에 있는 시트를 고른다(없으면 첫 시트)."""
    for ws in wb.worksheets:
        for row in range(FIRST_ROW, min(ws.max_row, FIRST_ROW + 4) + 1):
            if ws.cell(row=row, column=CODE_COL).value:
                return ws
    return wb.worksheets[0]


def _write_xlsx(os_label, results, out_path):
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        return None
    tpl_path = _find_template((os_label or "").upper())
    if not tpl_path:
        return None

    wb = openpyxl.load_workbook(tpl_path)
    ws = _pick_sheet(wb)
    row_of = {}
    for row in range(FIRST_ROW, ws.max_row + 1):
        code = ws.cell(row=row, column=CODE_COL).value
        if code is not None:
            row_of[str(code).strip()] = row

    red = Font(color="FFFF0000", bold=True)
    blue = Font(color="FF0070C0", bold=True)
    black = Font(color="FF000000")
    for r in results:
        row = row_of.get(str(r.get("code", "")))
        if not row:
            continue
        verdict = REPORT_STATUS.get(r.get("status", ""), r.get("status", ""))
        cell = ws.cell(row=row, column=RESULT_COL, value=verdict)
        cell.font = red if verdict == "취약" else blue if verdict == MANUAL_LABEL else black
        ws.cell(row=row, column=DETAIL_COL, value=" / ".join(r.get("evidence", []))[:32000])
        ws.cell(row=row, column=RES_COL, value=" / ".join(r.get("resources", []))[:32000])
    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------ 메인
def main():
    ap = argparse.ArgumentParser(description="클라우드 취약점 진단(GUI 없이 실행 → 엑셀 저장)")
    ap.add_argument("provider", choices=["aws", "azure", "gcp", "naver"], help="점검 대상 CSP")
    ap.add_argument("-o", "--output", help="저장 경로(확장자 제외 기본명). 기본: 현재 폴더에 cloud_<CSP>_<시각>")
    ap.add_argument("--json", dest="json_out", help="결과를 JSON 으로 저장(openpyxl 없는 곳에서 뽑아, PC 에서 make_report 로 xlsx 생성용)")
    ap.add_argument("--no-excel", action="store_true", help="xlsx 저장 생략(CSV 는 항상 저장)")
    ap.add_argument("--all-regions", action="store_true", help="AWS 전 리전 스캔(기본은 지정 리전만 → 빠름)")
    # 공통/개별 자격증명
    ap.add_argument("--access-key"); ap.add_argument("--secret-key")
    ap.add_argument("--session-token"); ap.add_argument("--profile")
    ap.add_argument("--region")
    ap.add_argument("--tenant-id"); ap.add_argument("--client-id")
    ap.add_argument("--client-secret"); ap.add_argument("--subscription-id")
    ap.add_argument("--sa-key"); ap.add_argument("--project")
    args = ap.parse_args()

    provider = args.provider.lower()
    if args.access_key or args.secret_key or args.client_secret:
        print("[!] 경고: 비밀키를 명령행 인자로 전달하면 ps/셸 히스토리/CI 로그에 노출됩니다. "
              "환경변수 또는 ambient 자격(CloudShell/역할/az login/ADC) 사용을 권장합니다.",
              file=sys.stderr)
    if args.all_regions:
        os.environ["CLOUD_SCAN_ALL_REGIONS"] = "1"
    if not cloud_check.available(provider):
        pkg = cloud_check.missing_package(provider)
        sys.exit(f"[!] {cloud_check.label(provider)} 진단 SDK 미설치 → pip install {pkg}")

    creds = collect_creds(provider, args)
    print(f"\n[*] {cloud_check.label(provider)} 진단 시작 …")
    try:
        report = cloud_check.run(provider, creds)
    except RuntimeError as e:
        sys.exit(f"[!] 진단 중단: {e}")
    except Exception as e:  # noqa: BLE001
        sys.exit(f"[!] 예기치 못한 오류({type(e).__name__}): {e}")

    results = report.get("results", [])
    host = report.get("host", "")
    os_label = report.get("os", provider.upper())

    # 콘솔 요약
    counts = {}
    print(f"\n===== 진단 결과: {host} ({os_label}) — {len(results)}항목 =====")
    for r in results:
        st = REPORT_STATUS.get(r.get("status", ""), r.get("status", ""))
        counts[st] = counts.get(st, 0) + 1
        mark = {"취약": "✗", "양호": "✓", "N/A": "-", MANUAL_LABEL: "?"}.get(st, " ")
        print(f"  {mark} [{r.get('code',''):<6}] {st:<8} {r.get('title','')}")
    print("  " + " | ".join(f"{k} {v}" for k, v in sorted(counts.items())))

    # JSON 저장(요청 시) — openpyxl 없는 클라우드 쉘에서 뽑아 PC 에서 make_report 로 xlsx 생성
    if args.json_out:
        import json as _json
        with open(args.json_out, "w", encoding="utf-8") as _f:
            _json.dump({"provider": provider, "host": host, "os": os_label, "results": results},
                       _f, ensure_ascii=False)
        print(f"[+] JSON 저장: {args.json_out}  (PC 에서: python make_report.py {provider} --result {args.json_out})")

    # 저장: CSV(항상) + 보고서 XLSX(openpyxl 있으면)
    paths = save_results(provider, os_label, host, results, args.output, want_xlsx=not args.no_excel)
    print()
    for pth in paths:
        print(f"[+] 저장: {pth}")
    if not any(p.endswith(".xlsx") for p in paths) and not args.no_excel:
        print("    (xlsx 보고서는 openpyxl 이 없어 생략 — CSV 를 열거나 pip install openpyxl 후 재실행)")
    print("[*] 완료")


if __name__ == "__main__":
    main()
