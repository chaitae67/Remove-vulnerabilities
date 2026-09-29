#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""서버(Linux/Windows) 진단 결과 → 공식 결과보고서 양식(다중서버) 채우기.

보고서_양식_{Linux,Windows}.xlsx 를 openpyxl 로 열어 표지/진단대상/상세(3-1)를 채우고,
2-2 요약과 2-1 그래프(3D막대·3D원형·레이더)는 템플릿 수식으로 자동 계산되게 한다.
리눅스는 2-2 HLOOKUP 수식과 그래프 시트 수식·차트를 명시적으로 복원/재구성한다.

make_report.py(CLI)가 servers 리스트를 넘겨 호출한다.
"""
import datetime

MANUAL_LABEL = "인터뷰 필요"
# 서버 보고서는 양호/취약/N/A 3분류 → 수동확인/인터뷰는 N/A(대상없음/판단보류)로 표기
RESULT_MAP = {
    "양호": "양호", "취약": "취약", "N/A": "N/A", "N/a": "N/A",
    "수동확인": "N/A", MANUAL_LABEL: "N/A", "": "N/A",
    "GOOD": "양호", "VULN": "취약", "NA": "N/A", "MAN": "N/A",
}

SPECS = {
    "linux": {
        "cover": "0. 표지", "target": "1. 진단 대상",
        "graph": "2-1. 요약결과(그래프)", "summary": "2-2. 요약 진단결과(Linux)",
        "detail": "3-1. 진단 결과(Linux)",
        "label": "Linux", "prefix": "U", "id_col": "C",
        "detail_first": 6, "detail_last": 72,
        "result_cols": ["F", "H", "J", "L"], "sum_cols": ["F", "G", "H", "I"],
        "servers": 4, "restore_formulas": True,
    },
    "windows": {
        "cover": "0. 표지", "target": "1. 진단 대상",
        "graph": "2-1. 요약결과(그래프)", "summary": "2-2. 요약 진단결과(Window)",
        "detail": "3-1. 진단 결과(window)",
        "label": "Windows", "prefix": "W", "id_col": "C",
        "detail_first": 6, "detail_last": 69,
        "result_cols": ["F", "H"], "sum_cols": ["F", "G"],
        "servers": 2, "restore_formulas": False,   # 윈도우 양식은 IFERROR 수식 완비 → 데이터만
    },
}
TARGET_FIRST_ROW = 5
COVER = {"docno": "L3", "author": "L4", "grade": "L5", "ver": "L6", "date": "B18"}


def _norm(v):
    return RESULT_MAP.get((v or "").strip(), RESULT_MAP.get((v or "").strip().upper(), "N/A"))


def _put(ws, addr, value):
    c = ws[addr]
    if type(c).__name__ == "MergedCell":
        raise RuntimeError(f"{ws.title}!{addr} 는 병합 셀입니다.")
    c.value = value


def _fill_cover(wb, spec, meta):
    ws = wb[spec["cover"]]
    _put(ws, COVER["docno"], meta.get("docno", "XXXXX-VA-2026XXX"))
    _put(ws, COVER["author"], meta.get("author", "취약점진단팀"))
    _put(ws, COVER["grade"], meta.get("grade", "Confidential"))
    _put(ws, COVER["ver"], meta.get("version", "ver 1.0"))
    _put(ws, COVER["date"], datetime.date.today())


def _fill_targets(wb, spec, servers):
    ws = wb[spec["target"]]
    for i, sv in enumerate(servers):
        r = TARGET_FIRST_ROW + i
        _put(ws, f"B{r}", i + 1)
        _put(ws, f"C{r}", (sv.get("host") or "").strip())
        _put(ws, f"D{r}", (sv.get("ip") or "-").strip())
        _put(ws, f"E{r}", (sv.get("osver") or "").strip())
        _put(ws, f"F{r}", (sv.get("role") or "-").strip())
    _put(ws, "B1", f"  ※ 진단 대상 리스트 - 서버 {len(servers)}대 ({spec['label']} {len(servers)}대)")


def _fill_detail(wb, spec, servers):
    ws = wb[spec["detail"]]
    first, last = spec["detail_first"], spec["detail_last"]
    row_of = {}
    for r in range(first, last + 1):
        v = ws[f"{spec['id_col']}{r}"].value
        if v:
            row_of[str(v).strip().upper()] = r
    for idx, sv in enumerate(servers):
        rc = spec["result_cols"][idx]
        ec = ws[f"{rc}1"].offset(0, 1).column_letter    # 판정 열 오른쪽 = 근거 열
        by_code = {(r.get("code", "") or "").strip().upper(): r for r in sv.get("results", [])}
        for code, r in row_of.items():
            x = by_code.get(code)
            if x is None:
                _put(ws, f"{rc}{r}", "N/A"); _put(ws, f"{ec}{r}", "점검 결과 없음")
                continue
            _put(ws, f"{rc}{r}", _norm(x.get("final") or x.get("status", "")))
            ev = " / ".join(x.get("evidence", []))
            note = (x.get("note") or "").strip()
            if note:
                ev = f"{ev}  [검증자: {note}]" if ev else f"[검증자: {note}]"
            _put(ws, f"{ec}{r}", ev)


def _restore_summary_formulas(wb, spec):
    ws = wb[spec["summary"]]
    first, last = spec["detail_first"], spec["detail_last"]
    det = spec["detail"]
    for r in range(first, last + 1):
        for c in spec["sum_cols"]:
            ws[f"{c}{r}"].value = (
                f"=HLOOKUP({c}$3,'{det}'!$F$3:$M$72,ROW(A{r-2}),FALSE)")


def _fix_graph_linux(wb, spec, n):
    from openpyxl.chart import BarChart3D, PieChart3D, RadarChart, Reference
    from openpyxl.chart.label import DataLabelList
    ws = wb[spec["graph"]]
    S_TARGET, S_SUM = spec["target"], spec["summary"]
    last = TARGET_FIRST_ROW + n - 1
    tgt = f"'{S_TARGET}'!$B${TARGET_FIRST_ROW}:$B${last}"
    sm = f"'{S_SUM}'!$F$74:$I$74"
    ws["C5"] = f"=AVERAGE({sm})"; ws["C6"] = f"=AVERAGE({sm})"
    ws["D5"] = "=D6"; ws["D6"] = f"=COUNTA({tgt})"
    ws["D17"] = f"=COUNTA({tgt})"
    ws["D18"] = f'=COUNTIF({sm},">=0.85")'
    ws["D19"] = "=D17-(D18+D20)"
    ws["D20"] = f'=COUNTIF({sm},"<0.7")'
    for i, r in enumerate(range(72, 77)):
        src = [6, 19, 39, 69, 70][i]
        ws[f"B{r}"] = f"='{S_SUM}'!B{src}"
        ws[f"C{r}"] = f"='{S_SUM}'!K{src}"
        ws[f"D{r}"] = "=$C$6"
    for r in range(36, 51):
        if r - 35 <= n:
            ws[f"W{r}"] = f"=HLOOKUP($V{r},'{S_SUM}'!$F$3:$I$53,2,0)"
            ws[f"X{r}"] = f"=HLOOKUP($V{r},'{S_SUM}'!$F$3:$I$74,72,0)"
        else:
            ws[f"W{r}"] = None; ws[f"X{r}"] = None
    ws._charts = []
    bar = BarChart3D(); bar.type = "col"; bar.legend = None
    bar.add_data(Reference(ws, min_col=3, min_row=4, max_row=6), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=5, max_row=6))
    bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True
    bar.width, bar.height = 12.6, 8.1
    ws.add_chart(bar, "E3")
    pie = PieChart3D()
    pie.add_data(Reference(ws, min_col=3, min_row=18, max_row=20), titles_from_data=False)
    pie.set_categories(Reference(ws, min_col=2, min_row=18, max_row=20))
    pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
    pie.width, pie.height = 12.6, 7.1
    ws.add_chart(pie, "E31")
    radar = RadarChart(); radar.type = "marker"
    radar.add_data(Reference(ws, min_col=3, max_col=4, min_row=71, max_row=76), titles_from_data=True)
    radar.set_categories(Reference(ws, min_col=2, min_row=72, max_row=76))
    radar.y_axis.scaling.min, radar.y_axis.scaling.max = 0, 1
    radar.width, radar.height = 13.6, 9.3
    ws.add_chart(radar, "G71")


# ---------------- 진입점 ----------------
def fill_report(os_kind, servers, template_path, out_path, meta=None):
    """서버 여러 대 결과를 계열 양식(다중서버)에 채워 저장.

    os_kind : 'linux' | 'windows'
    servers : [{host, ip, osver, role, results:[{code,status,evidence,...}]}...]
              (단일 dict / results 리스트도 허용)
    """
    from openpyxl import load_workbook
    spec = SPECS[os_kind]
    meta = meta or {}

    if isinstance(servers, dict):
        servers = [servers]
    elif servers and isinstance(servers[0], dict) and "results" not in servers[0] and "code" in servers[0]:
        servers = [{"results": servers}]
    servers = [s for s in servers if s.get("results")][: spec["servers"]]
    if not servers:
        raise ValueError("진단 결과(servers)가 비어 있습니다.")

    wb = load_workbook(template_path)
    wb._external_links = []
    _fill_cover(wb, spec, meta)
    _fill_targets(wb, spec, servers)
    _fill_detail(wb, spec, servers)
    if spec["restore_formulas"]:
        _restore_summary_formulas(wb, spec)
        _fix_graph_linux(wb, spec, len(servers))
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path
